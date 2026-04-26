import os
# Memory optimizations for Render free tier
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"
os.environ["PYTORCH_CPU_ALLOC_CONF"] = "max_split_size_mb:128"

import torch
torch.set_num_threads(1)
torch.set_num_interop_threads(1)
from flask import Flask, render_template, request, send_file
from ultralytics import YOLO
import os
import cv2
from fpdf import FPDF
from datetime import datetime
import csv, io, json

app = Flask(__name__)
model = YOLO("best (3).pt")

UPLOAD_FOLDER = "static/uploads"
REPORTS_FOLDER = "static/reports"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(REPORTS_FOLDER, exist_ok=True)

COST_RATES = {
    "Pothole": {"base": 1500, "per_pixel": 0.05},
    "Manhole": {"base": 2000, "per_pixel": 0.03},
    "Longitudinal Crack": {"base": 800, "per_pixel": 0.02},
    "Transverse Crack": {"base": 600, "per_pixel": 0.02},
    "Alligator Crack": {"base": 1200, "per_pixel": 0.04}
}

SEVERITY_MULTIPLIERS = {
    "critical": 2.5,
    "high": 1.8,
    "moderate": 1.2,
    "low": 0.8
}

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

CLASS_NAMES = {
    0: "Longitudinal Crack",
    1: "Transverse Crack",
    2: "Alligator Crack",
    3: "Manhole"
}


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        file = request.files["file"]
        conf_threshold = float(request.form.get("conf_threshold", 0.25))

        if file:
            filepath = os.path.join(app.config["UPLOAD_FOLDER"], file.filename)
            file.save(filepath)
            ext = file.filename.split(".")[-1].lower()

            if ext in ["jpg", "jpeg", "png"]:
                results = model.predict(source=filepath, conf=conf_threshold)
                annotated = results[0].plot()
                output_path = os.path.join(app.config["UPLOAD_FOLDER"], "output.jpg")
                cv2.imwrite(output_path, annotated)

                detections = []
                cost_breakdown = []
                total_cost = 0

                for box in results[0].boxes:
                    cls = int(box.cls[0])
                    conf = float(box.conf[0])
                    label = CLASS_NAMES.get(cls, "Unknown")
                    detections.append(f"{label} ({conf:.2f})")

                    x1, y1, x2, y2 = box.xyxy[0].tolist()
                    area = (x2 - x1) * (y2 - y1)

                    if conf > 0.8:
                        severity = "critical"
                    elif conf > 0.6:
                        severity = "high"
                    elif conf > 0.4:
                        severity = "moderate"
                    else:
                        severity = "low"

                    rates = COST_RATES.get(label, {"base": 1000, "per_pixel": 0.02})
                    cost = (rates["base"] + area * rates["per_pixel"]) * SEVERITY_MULTIPLIERS[severity]
                    cost = round(cost, 2)
                    total_cost += cost

                    cost_breakdown.append({
                        "type": label,
                        "confidence": conf,
                        "severity": severity,
                        "area_pixels": int(area),
                        "cost": cost
                    })

                pdf_path = generate_pdf_report(filepath, output_path, detections, cost_breakdown, total_cost)
                save_to_governance_excel(file.filename, detections, cost_breakdown, total_cost)

                return render_template("index.html",
                                       input_image=filepath,
                                       output_image=output_path,
                                       detections=detections,
                                       cost_breakdown=cost_breakdown,
                                       total_cost=round(total_cost, 2),
                                       pdf_report=pdf_path)

            elif ext in ["mp4", "avi", "mov"]:
                print(f"[DEBUG] ========== VIDEO PROCESSING START ==========", flush=True)
                print(f"[DEBUG] File: {file.filename}, Ext: {ext}", flush=True)
                print(f"[DEBUG] Filepath: {filepath}", flush=True)
                print(f"[DEBUG] File exists: {os.path.exists(filepath)}", flush=True)

                try:
                    # Process video with YOLO
                    print(f"[DEBUG] Starting YOLO prediction...")
                    results = model.predict(
                        source=filepath,
                        conf=conf_threshold,
                        save=True,
                        project="static",
                        name="output_video",
                        exist_ok=True
                    )
                    print(f"[DEBUG] YOLO prediction completed!")

                    # Get YOLO save directory
                    save_dir = str(results[0].save_dir)
                    print(f"[DEBUG] YOLO save_dir: {save_dir}")
                except Exception as e:
                    print(f"[DEBUG] ERROR in YOLO prediction: {e}")
                    import traceback
                    traceback.print_exc()

                # Find YOLO output file
                base_name = os.path.splitext(file.filename)[0]  # Use original filename, not filepath
                yolo_output = None
                for candidate_ext in [".mp4", ".avi", ".mov"]:
                    candidate = os.path.join(save_dir, base_name + candidate_ext)
                    print(f"[DEBUG] Checking: {candidate} - Exists: {os.path.exists(candidate)}")
                    if os.path.exists(candidate):
                        yolo_output = candidate
                        break

                # Setup output path
                output_dir = os.path.join("static", "output_video")
                os.makedirs(output_dir, exist_ok=True)
                output_file = base_name + "_annotated.mp4"
                output_path = os.path.join(output_dir, output_file)
                print(f"[DEBUG] Output path: {output_path}")

                # Re-encode video
                if yolo_output and os.path.exists(yolo_output):
                    print(f"[DEBUG] Re-encoding from: {yolo_output}")
                    cap = cv2.VideoCapture(yolo_output)
                    fps = cap.get(cv2.CAP_PROP_FPS) or 25
                    w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
                    h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
                    print(f"[DEBUG] Video properties: {w}x{h} @ {fps}fps")

                    # Try H.264 codec first (best browser compatibility), fallback to mp4v
                    fourcc = cv2.VideoWriter_fourcc(*"avc1")
                    writer = cv2.VideoWriter(output_path, fourcc, fps, (w, h))
                    if not writer.isOpened():
                        print(f"[DEBUG] avc1 failed, trying mp4v...")
                        fourcc = cv2.VideoWriter_fourcc(*"mp4v")
                        writer = cv2.VideoWriter(output_path, fourcc, fps, (w, h))
                    if not writer.isOpened():
                        print(f"[DEBUG] mp4v failed, trying XVID...")
                        fourcc = cv2.VideoWriter_fourcc(*"XVID")
                        writer = cv2.VideoWriter(output_path, fourcc, fps, (w, h))

                    if writer.isOpened():
                        print(f"[DEBUG] VideoWriter opened successfully with codec {fourcc}")
                    else:
                        print(f"[DEBUG] ERROR: VideoWriter failed to open!")

                    frame_count = 0
                    while True:
                        ret, frame = cap.read()
                        if not ret:
                            break
                        writer.write(frame)
                        frame_count += 1

                    cap.release()
                    writer.release()
                    print(f"[DEBUG] Wrote {frame_count} frames to {output_path}")
                    print(f"[DEBUG] Output file exists: {os.path.exists(output_path)}, Size: {os.path.getsize(output_path) if os.path.exists(output_path) else 0}")
                else:
                    print(f"[DEBUG] YOLO output not found!")

                # URL paths for template
                input_video_url = "static/uploads/" + file.filename
                output_video_url = "static/output_video/" + output_file
                print(f"[DEBUG] URLs - Input: {input_video_url}, Output: {output_video_url}")

                return render_template(
                    "index.html",
                    input_video=input_video_url,
                    output_video=output_video_url,
                    conf_threshold=conf_threshold
                )

    return render_template("index.html")


@app.route("/download")
def download():
    path = os.path.join(UPLOAD_FOLDER, "output.jpg")
    return send_file(path, as_attachment=True)


@app.route("/download_report")
def download_report():
    pdf_path = request.args.get("pdf", "")
    if pdf_path and os.path.exists(pdf_path):
        return send_file(pdf_path, as_attachment=True,
                         download_name=f"road_damage_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
    return "Report not found", 404


def generate_pdf_report(input_img, output_img, detections, cost_breakdown, total_cost):
    pdf_path = os.path.join(REPORTS_FOLDER, f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    PAGE_W = pdf.w - pdf.l_margin - pdf.r_margin
    IMG_W = PAGE_W * 0.60
    IMG_X = pdf.l_margin + (PAGE_W - IMG_W) / 2
    GAP = 8

    pdf.set_font("Arial", "B", 24)
    pdf.cell(0, 15, "Road Damage Detection Report", ln=True, align="C")

    pdf.set_font("Arial", "", 11)
    pdf.cell(0, 8, f"Generated: {datetime.now().strftime('%B %d, %Y at %H:%M')}", ln=True, align="C")
    pdf.ln(5)

    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Detected Road Damage", ln=True, align="C")
    pdf.ln(2)

    img_y = pdf.get_y()

    try:
        _img = cv2.imread(output_img)
        if _img is not None:
            px_h, px_w = _img.shape[:2]
            IMG_H = IMG_W * (px_h / px_w)
        else:
            IMG_H = IMG_W * 0.75
    except Exception:
        IMG_H = IMG_W * 0.75

    try:
        pdf.image(output_img, x=IMG_X, y=img_y, w=IMG_W)
    except Exception:
        pass

    pdf.set_y(img_y + IMG_H + GAP)

    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Detection Summary", ln=True)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(5)

    pdf.set_font("Arial", "B", 11)
    pdf.cell(60, 8, "Damage Type", 1)
    pdf.cell(40, 8, "Count", 1, align="C")
    pdf.cell(50, 8, "Avg Confidence", 1, align="C")
    pdf.cell(40, 8, "Status", 1, align="C", ln=True)

    type_counts = {}
    type_conf = {}
    for item in cost_breakdown:
        t = item["type"]
        type_counts[t] = type_counts.get(t, 0) + 1
        type_conf[t] = type_conf.get(t, []) + [item["confidence"]]

    pdf.set_font("Arial", "", 11)
    for dtype, count in type_counts.items():
        avg_conf = sum(type_conf[dtype]) / len(type_conf[dtype])
        pdf.cell(60, 8, dtype, 1)
        pdf.cell(40, 8, str(count), 1, align="C")
        pdf.cell(50, 8, f"{avg_conf:.2%}", 1, align="C")

        if count >= 5:
            status_text = "Severe"
            pdf.set_text_color(200, 0, 0)
        elif count >= 2:
            status_text = "Moderate"
            pdf.set_text_color(230, 150, 0)
        else:
            status_text = "Minor"
            pdf.set_text_color(0, 150, 0)

        pdf.cell(40, 8, status_text, 1, align="C", ln=True)
        pdf.set_text_color(0, 0, 0)

    pdf.ln(10)

    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "AI-Powered Repair Cost Estimation", ln=True)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(5)

    pdf.set_font("Arial", "B", 10)
    pdf.cell(50, 8, "Type", 1)
    pdf.cell(35, 8, "Severity", 1, align="C")
    pdf.cell(35, 8, "Area (px)", 1, align="C")
    pdf.cell(35, 8, "Cost (Rs.)", 1, align="C", ln=True)

    pdf.set_font("Arial", "", 10)
    for item in cost_breakdown:
        pdf.cell(50, 7, item["type"][:25], 1)
        pdf.cell(35, 7, item["severity"].upper(), 1, align="C")
        pdf.cell(35, 7, f"{item['area_pixels']:,}", 1, align="C")
        pdf.cell(35, 7, f"Rs.{item['cost']:,.2f}", 1, align="C", ln=True)

    pdf.ln(3)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(120, 10, "ESTIMATED TOTAL COST:", 0, align="R")
    pdf.cell(0, 10, f"Rs.{total_cost:,.2f}", 0, ln=True)
    pdf.ln(15)

    pdf.set_font("Arial", "I", 9)
    pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(
        0, 5,
        "Note: Cost estimates are based on bounding box area and detection confidence. "
        "Actual repair costs may vary based on material prices, labor rates, and site conditions."
    )

    pdf.output(pdf_path)
    return pdf_path


def save_to_governance_excel(filename, detections, cost_breakdown, total_cost):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return

    excel_path = os.path.join(REPORTS_FOLDER, "governance_records.xlsx")
    timestamp = datetime.now()

    if os.path.exists(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Road Damage Records"
        headers = ["Date", "Time", "Filename", "Total Detections", "Total Cost (Rs.)",
                   "Detection #", "Damage Type", "Confidence", "Severity", "Area (px)", "Cost (Rs.)", "Status"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start_row = ws.max_row + 1

    if cost_breakdown:
        num_detections = len(cost_breakdown)

        for i, item in enumerate(cost_breakdown):
            row = [
                timestamp.strftime("%Y-%m-%d") if i == 0 else "",
                timestamp.strftime("%H:%M:%S") if i == 0 else "",
                filename if i == 0 else "",
                len(detections) if i == 0 else "",
                f"Rs.{total_cost:,.2f}" if i == 0 else "",
                i + 1,
                item['type'],
                f"{item['confidence']:.2%}",
                item["severity"].upper(),
                f"{item['area_pixels']:,}",
                f"Rs.{item['cost']:,.2f}",
                "Active" if item["confidence"] > 0.5 else "Review"
            ]
            ws.append(row)

        if num_detections > 1:
            for col in range(1, 6):
                ws.merge_cells(start_row=start_row, start_column=col,
                               end_row=start_row + num_detections - 1, end_column=col)
                ws.cell(row=start_row, column=col).alignment = Alignment(
                    horizontal="center", vertical="center")

        for row in range(start_row, start_row + num_detections):
            ws.row_dimensions[row].height = 20
    else:
        row = [
            timestamp.strftime("%Y-%m-%d"),
            timestamp.strftime("%H:%M:%S"),
            filename,
            0,
            "Rs.0.00",
            "-",
            "No damage detected",
            "-",
            "-",
            "-",
            "-",
            "Clean"
        ]
        ws.append(row)

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

    wb.save(excel_path)


@app.route("/download/governance_excel")
def download_governance_excel():
    excel_path = os.path.join(REPORTS_FOLDER, "governance_records.xlsx")
    if os.path.exists(excel_path):
        return send_file(excel_path, as_attachment=True,
                         download_name=f"road_damage_governance_{datetime.now().strftime('%Y%m%d')}.xlsx")
    return "No governance records found yet. Upload an image first.", 404


@app.route("/clear_governance_records")
def clear_governance_records():
    excel_path = os.path.join(REPORTS_FOLDER, "governance_records.xlsx")
    if os.path.exists(excel_path):
        os.remove(excel_path)
        return "All governance records deleted successfully!", 200
    return "No records found to delete.", 404


@app.route("/download/csv")
def download_csv():
    detections = app.config.get('LAST_DETECTIONS', [])
    filename = app.config.get('LAST_FILENAME', 'unknown')
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['#', 'Label', 'Confidence', 'Severity', 'X1', 'Y1', 'X2', 'Y2', 'Timestamp', 'Source File'])
    for i, d in enumerate(detections, 1):
        writer.writerow([i, d['label'], d['conf'], d['severity'],
                         d['x1'], d['y1'], d['x2'], d['y2'],
                         datetime.now().strftime('%Y-%m-%d %H:%M:%S'), filename])
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode()),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'road_damage_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
    )


@app.route("/download/pdf")
def download_pdf():
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
    except ImportError:
        return "Run: pip install reportlab", 500

    detections = app.config.get('LAST_DETECTIONS', [])
    filename = app.config.get('LAST_FILENAME', 'unknown')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Road Damage Detection Report", styles['Title']))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph(f"<b>Source File:</b> {filename}", styles['Normal']))
    story.append(Paragraph(f"<b>Date:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    story.append(Paragraph(f"<b>Total Detections:</b> {len(detections)}", styles['Normal']))
    story.append(Spacer(1, 0.6 * cm))

    data = [['#', 'Damage Type', 'Confidence', 'Severity', 'Bbox (x1,y1,x2,y2)']]
    for i, d in enumerate(detections, 1):
        data.append([str(i), d['label'], f"{d['conf']:.2f}", d['severity'],
                     f"({d['x1']},{d['y1']},{d['x2']},{d['y2']})"])

    table = Table(data, colWidths=[1 * cm, 5 * cm, 3 * cm, 3 * cm, 5 * cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f5a623')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                     download_name=f'road_damage_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf')


if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
